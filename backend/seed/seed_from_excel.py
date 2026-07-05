"""
Reads seed/thoughtify_seed_data.xlsx and inserts the data into your Supabase Postgres
database (users + profiles + categories + posts), using the same models/DB connection
as the FastAPI app (reads DATABASE_URL from backend/.env).

Usage:
    cd backend
    python seed/seed_from_excel.py                       # uses seed/thoughtify_seed_data.xlsx
    python seed/seed_from_excel.py path/to/other.xlsx     # or point at your own file

Excel format expected:

Sheet "users":
    username | email | password | first_name | last_name | avatar_url | phone | dob | address

Sheet "posts":
    title | content | category | image_url | author_username | status | section | main_post

- `avatar_url` / `image_url` can be ANY public image URL (Cloudinary, Pravatar, Picsum, etc.)
- `author_username` in the posts sheet must match a `username` from the users sheet
- `status` must be "draft" or "publish"; `section` must be "recent", "popular", or "trending"
- `main_post` should be TRUE/FALSE (only one post should really be TRUE, it's the homepage hero)

Re-running is safe: existing users/categories/posts (matched by username / name / title)
are skipped rather than duplicated.
"""
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))  # allow `import app...`

from app.database import SessionLocal, Base, engine  # noqa: E402
from app import models  # noqa: E402
from app.auth import hash_password  # noqa: E402
from app.utils import unique_slug  # noqa: E402


def load_workbook(path: str):
    wb = openpyxl.load_workbook(path, data_only=True)
    return wb


def rows(ws):
    """Yield each data row (skipping the header) as a list of cell values."""
    header_skipped = False
    for row in ws.iter_rows(values_only=True):
        if not header_skipped:
            header_skipped = True
            continue
        if row[0] is None:  # skip blank trailing rows
            continue
        yield row


def seed_users(db, ws):
    created, skipped = 0, 0
    for username, email, password, first_name, last_name, avatar_url, phone, dob, address in rows(ws):
        existing = db.query(models.User).filter(models.User.username == username).first()
        if existing:
            skipped += 1
            continue

        user = models.User(
            username=username,
            email=email,
            hashed_password=hash_password(str(password)),
            first_name=first_name or "",
            last_name=last_name or "",
        )
        db.add(user)
        db.flush()  # get user.id

        profile = models.UserProfile(
            user_id=user.id,
            image=avatar_url or None,
            phone=str(phone) if phone else None,
            dob=str(dob) if dob else None,
            address=address or None,
        )
        db.add(profile)
        created += 1

    db.commit()
    print(f"Users: {created} created, {skipped} already existed")


def get_or_create_category(db, name):
    category = db.query(models.Category).filter(models.Category.name == name).first()
    if category:
        return category
    category = models.Category(name=name, slug=unique_slug(db, models.Category, name))
    db.add(category)
    db.flush()
    return category


def seed_posts(db, ws):
    created, skipped = 0, 0
    for title, content, category_name, image_url, author_username, status, section, main_post in rows(ws):
        if db.query(models.Post).filter(models.Post.title == title).first():
            skipped += 1
            continue

        author = db.query(models.User).filter(models.User.username == author_username).first()
        if not author:
            print(f"  ! Skipping '{title}': no user with username '{author_username}'")
            continue

        category = get_or_create_category(db, category_name or "General")

        post = models.Post(
            title=title,
            content=content,
            image=image_url or None,
            category_id=category.id,
            blog_slug=unique_slug(db, models.Post, title),
            status=(status or "draft").strip().lower(),
            section=(section or "recent").strip().lower(),
            main_post=bool(main_post) if main_post is not None else False,
            author_id=author.id,
        )
        db.add(post)
        created += 1

    db.commit()
    print(f"Posts: {created} created, {skipped} already existed")


def main():
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else str(Path(__file__).parent / "thoughtify_seed_data.xlsx")
    if not Path(xlsx_path).exists():
        print(f"File not found: {xlsx_path}")
        print("Generate the sample first with: python seed/generate_sample_excel.py")
        sys.exit(1)

    print(f"Loading {xlsx_path} ...")
    wb = load_workbook(xlsx_path)

    # Make sure tables exist (harmless if they already do — same call main.py makes on boot).
    Base.metadata.create_all(bind=engine)

    db = SessionLocal()
    try:
        seed_users(db, wb["users"])
        seed_posts(db, wb["posts"])
    finally:
        db.close()

    print("Done. Check Supabase Table Editor to confirm.")


if __name__ == "__main__":
    main()